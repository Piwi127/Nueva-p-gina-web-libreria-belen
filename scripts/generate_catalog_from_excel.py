import concurrent.futures
import html
import json
import re
import unicodedata
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from urllib.error import HTTPError, URLError
from urllib.parse import quote, urlsplit, urlunsplit
from urllib.request import Request, urlopen

import difflib
import openpyxl


ROOT = Path(__file__).resolve().parents[1]
DATA_XLSX = ROOT / "data" / "articulos_con_precios_para_venta.xlsx"
PRODUCTS_JS = ROOT / "data" / "products.js"
IMG_PRODUCTS_DIR = ROOT / "img" / "products"
USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"


def strip_accents(text: str) -> str:
    normalized = unicodedata.normalize("NFD", text)
    return "".join(ch for ch in normalized if unicodedata.category(ch) != "Mn")


def normalize_text(text: str) -> str:
    text = strip_accents(str(text or "")).lower()
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def normalize_for_match(text: str) -> str:
    base = normalize_text(text)
    base = re.sub(r"\bx\s*1\s*und(?:\b|$)", " ", base)
    base = re.sub(r"\bx\s*1\s*unidad(?:es)?(?:\b|$)", " ", base)
    base = re.sub(r"\s+", " ", base).strip()
    return base


def slugify(text: str) -> str:
    slug = normalize_text(text).replace(" ", "-")
    slug = re.sub(r"-{2,}", "-", slug).strip("-")
    return slug or "producto"


def safe_float(value) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    text = text.replace("S/", "").replace("s/", "").replace(",", ".")
    text = re.sub(r"[^0-9.]", "", text)
    if text.count(".") > 1:
        parts = text.split(".")
        text = "".join(parts[:-1]) + "." + parts[-1]
    try:
        return float(text)
    except ValueError:
        return None


def clean_html_to_text(value: str) -> str:
    if not value:
        return ""
    text = re.sub(r"<br\s*/?>", "\n", value, flags=re.I)
    text = re.sub(r"</p>", "\n", text, flags=re.I)
    text = re.sub(r"<[^>]+>", " ", text)
    text = html.unescape(text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def first_sentences(text: str, count: int = 2) -> str:
    text = re.sub(r"\s+", " ", str(text or "")).strip()
    if not text:
        return ""
    parts = re.split(r"(?<=[.!?])\s+", text)
    parts = [p.strip() for p in parts if p.strip()]
    if not parts:
        return text
    joined = " ".join(parts[:count]).strip()
    return joined


def sanitize_url(raw: str) -> str:
    value = str(raw or "").strip()
    if not value:
        return ""
    try:
        parsed = urlsplit(value)
    except ValueError:
        return ""
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        return ""

    path = quote(parsed.path or "/", safe="/%")
    query = quote(parsed.query, safe="=&%")
    return urlunsplit((parsed.scheme, parsed.netloc, path, query, parsed.fragment))


def load_existing_products(path: Path) -> Dict[str, dict]:
    if not path.exists():
        return {}
    text = path.read_text(encoding="utf-8", errors="ignore")
    match = re.search(r"const PRODUCTS = (\{.*?\});\s*const OFFERS =", text, flags=re.S)
    if not match:
        return {}
    data = json.loads(match.group(1))
    return data


def parse_sheet_rows() -> List[dict]:
    wb = openpyxl.load_workbook(DATA_XLSX, data_only=True)
    ws_data = wb["NUEVA DATA"]
    ws_base = wb["NUEVA BASE"]

    base_map = {}
    for row in ws_base.iter_rows(min_row=2, values_only=True):
        code = str(row[0] or "").strip()
        name = str(row[1] or "").strip()
        price = safe_float(row[2])
        if not name:
            continue
        key = normalize_for_match(name)
        base_map[key] = {"code": code, "name": name, "price": price}

    rows: List[dict] = []
    seen = set()
    for i, row in enumerate(ws_data.iter_rows(min_row=2, values_only=True), start=2):
        raw_name = str(row[1] or "").strip()
        if not raw_name:
            continue

        key = normalize_for_match(raw_name)
        base = base_map.get(key, {})
        display_name = base.get("name") or raw_name
        code = base.get("code", "")

        price = (
            safe_float(row[7])
            or safe_float(row[8])
            or base.get("price")
            or safe_float(row[5])
            or safe_float(row[2])
            or 0.0
        )
        link = sanitize_url(row[6])
        if key in seen:
            continue
        seen.add(key)
        rows.append(
            {
                "excel_row": i,
                "code": code,
                "title": display_name,
                "price": float(price or 0.0),
                "link": link,
            }
        )
    return rows


def extract_brand_from_title(title: str) -> str:
    text = normalize_text(title)
    patterns = [
        "faber castell",
        "artesco",
        "pilot",
        "ove",
        "vinifan",
        "stabilo",
        "justus",
        "standford",
        "luxor",
        "layconsa",
        "david",
        "graphos",
        "chamex",
        "evaflex",
        "dragon",
        "alpha",
        "yimi",
    ]
    for brand in patterns:
        if brand in text:
            return " ".join(w.capitalize() for w in brand.split())
    return ""


def categorize_title(title: str, external_categories: Optional[List[str]] = None) -> str:
    t = normalize_text(title)
    ctext = normalize_text(" ".join(external_categories or []))
    source = f"{t} {ctext}".strip()

    if re.search(r"\b(cuaderno|block|cuadernillo|anillad|empastad|libreta)\b", source):
        return "cuadernos"
    if re.search(
        r"\b(cartulina|papel bond|papel lustre|papel crepe|papelografo|microporoso|papel fotocopia|papel autocopia|papel seda|papel kraft|papel carbon|resma|lustre|forro|mica)\b",
        source,
    ):
        return "papeleria"
    if re.search(
        r"\b(acuarela|tempera|pintura|oleo|pincel|crayon|plastilina|foamy|eva|origami|escarcha|glitter|lentejuela|palito)\b",
        source,
    ):
        return "arte"
    if re.search(
        r"\b(borrador|tajador|sacapuntas|regla|transportador|compas|pegamento|goma|cola|silicona|tijera|cartuchera|estuche|colores|corrector|limpiatipos)\b",
        source,
    ):
        return "utiles"
    if re.search(r"\b(boligrafo|lapicero|lapiz|plumon|marcador|resaltador|tinta)\b", source):
        return "escritura"
    if re.search(r"\b(archivador|folder|carpeta|separador|organizador|clip|chinche|chinches|alfiler|alfileres|broche|sobre|mica|forro)\b", source):
        return "organizacion"
    if re.search(r"\b(perforador|grapadora|engrap|grapa|grapas|tampon|sello|cinta adhesiva|dispensador)\b", source):
        return "oficina"
    if re.search(r"\b(usb|memoria|mouse|teclado|audifono|cd|dvd|laptop|bateria|cargador|toner)\b", source):
        return "tecnologia"
    if re.search(r"\b(regalo|detalle|juguete)\b", source):
        return "regalos"
    return "otros"


def build_usage(title: str, category: str) -> List[str]:
    t = normalize_text(title)
    if "cuaderno" in t or "block" in t:
        return [
            "Perfecto para clases, tareas y organizacion de apuntes diarios.",
            "Facilita la toma de notas, ejercicios y trabajos escolares.",
            "Complementa con separadores y forros para mayor orden y proteccion.",
        ]
    if "acuarela" in t or "tempera" in t or "pintura" in t:
        return [
            "Utiliza el producto para trabajos de arte y actividades escolares.",
            "Combina con pinceles y cartulina para obtener mejores acabados.",
            "Guarda en lugar seco para conservar su calidad por mas tiempo.",
        ]
    if "boligrafo" in t or "lapiz" in t or "plumon" in t or "resaltador" in t:
        return [
            "Ideal para escritura diaria en colegio, universidad y oficina.",
            "Permite tomar apuntes, subrayar ideas y completar tareas con precision.",
            "Manten la tapa cerrada o guardalo en estuche para mayor duracion.",
        ]
    if re.search(r"\b(pegamento|cola|silicona|goma|adhesivo)\b", t):
        return [
            "Apto para unir papel, cartulina y materiales ligeros en manualidades.",
            "Aplica una capa uniforme para una fijacion limpia y resistente.",
            "Espera unos minutos de secado antes de manipular la superficie.",
        ]
    if category == "papeleria":
        return [
            "Util para impresiones, recortes, forrado y trabajos escolares.",
            "Compatible con actividades de aula, oficina y manualidades.",
            "Conserva el material en superficie plana para evitar dobleces.",
        ]
    if category == "organizacion":
        return [
            "Permite clasificar y proteger documentos de uso frecuente.",
            "Facilita el archivo por temas, cursos o proyectos de trabajo.",
            "Ideal para mantener orden en escritorio, mochila o archivador.",
        ]
    if category == "oficina":
        return [
            "Pensado para tareas administrativas y uso diario en escritorio.",
            "Ayuda a ordenar, perforar, fijar o gestionar documentos de oficina.",
            "Recomendado para estudio, negocio y trabajo remoto.",
        ]
    return [
        "Producto recomendado para actividades escolares y de oficina.",
        "Facil de usar en tareas diarias, proyectos y organizacion personal.",
        "Conservalo en lugar seco para mantener su rendimiento.",
    ]


def extract_specs(title: str) -> List[str]:
    lower = normalize_text(title)
    specs = []

    pack = re.search(r"\bx\s*(\d+)\s*(und|un|hojas|gr|g|mm|ml|yd)\b", lower)
    if pack:
        qty, unit = pack.groups()
        unit_map = {
            "und": "unidades",
            "un": "unidades",
            "hojas": "hojas",
            "gr": "gramos",
            "g": "gramos",
            "mm": "mm",
            "ml": "ml",
            "yd": "yardas",
        }
        specs.append(f"Presentacion: x {qty} {unit_map.get(unit, unit)}.")

    size = re.search(r"\b(a3|a4|a5|oficio|carta)\b", lower)
    if size:
        specs.append(f"Formato: {size.group(1).upper()}.")

    dims = re.search(r"\b(\d+(?:\.\d+)?)\s*x\s*(\d+(?:\.\d+)?)\s*cm\b", lower)
    if dims:
        specs.append(f"Medida: {dims.group(1)} x {dims.group(2)} cm.")

    if "gramos" in lower or re.search(r"\b\d+\s*g\b", lower):
        grams = re.search(r"\b(\d+(?:\.\d+)?)\s*(?:gramos|g)\b", lower)
        if grams:
            specs.append(f"Peso o gramaje: {grams.group(1)} g.")

    return specs


def extract_colors(title: str) -> List[str]:
    color_map = {
        "negro": "Negro",
        "blanco": "Blanco",
        "azul": "Azul",
        "rojo": "Rojo",
        "verde": "Verde",
        "amarillo": "Amarillo",
        "naranja": "Naranja",
        "rosado": "Rosado",
        "marron": "Marron",
        "morado": "Morado",
        "lila": "Lila",
        "celeste": "Celeste",
        "turquesa": "Turquesa",
        "fucsia": "Fucsia",
        "gris": "Gris",
        "dorado": "Dorado",
    }
    text = normalize_text(title)
    colors = [label for key, label in color_map.items() if re.search(rf"\b{re.escape(key)}\b", text)]
    return colors[:4] or ["Clasico"]


def extract_sizes(title: str) -> List[str]:
    text = normalize_text(title)
    sizes = []
    for token in re.findall(r"\b(a3|a4|a5|oficio|carta)\b", text):
        sizes.append(token.upper())
    for token in re.findall(r"\b\d+(?:\.\d+)?\s*(?:mm|ml|g|gr)\b", text):
        cleaned = token.upper().replace("GR", "G")
        sizes.append(cleaned)
    unique = []
    seen = set()
    for item in sizes:
        if item in seen:
            continue
        seen.add(item)
        unique.append(item)
    return unique[:4] or ["Unico"]


def parse_next_data_from_html(html_text: str) -> Optional[dict]:
    match = re.search(
        r'<script id="__NEXT_DATA__" type="application/json">(.*?)</script>',
        html_text,
        flags=re.S,
    )
    if not match:
        return None
    try:
        payload = json.loads(match.group(1))
    except json.JSONDecodeError:
        return None

    page_props = payload.get("props", {}).get("pageProps", {})
    product = page_props.get("product", {}) or {}
    seo = page_props.get("seo", {}) or {}
    brand = product.get("brand") or {}
    categories = product.get("categories") or []
    category_names = [c.get("name", "").strip() for c in categories if isinstance(c, dict) and c.get("name")]

    return {
        "name": (product.get("name") or "").strip(),
        "description_html": product.get("description") or "",
        "short_description_html": product.get("short_description") or "",
        "meta_description": seo.get("description") or product.get("meta_description") or "",
        "brand": (brand.get("name") if isinstance(brand, dict) else "") or "",
        "categories": category_names,
        "price": safe_float(product.get("price")),
        "image_url": product.get("image_url") or "",
    }


def fetch_product_info(url: str) -> Tuple[str, Optional[dict], Optional[str]]:
    if not url:
        return url, None, "missing_url"
    try:
        req = Request(url, headers={"User-Agent": USER_AGENT})
        with urlopen(req, timeout=20) as resp:
            raw = resp.read()
        html_text = raw.decode("utf-8", errors="ignore")
        data = parse_next_data_from_html(html_text)
        if not data:
            return url, None, "no_next_data"
        return url, data, None
    except HTTPError as error:
        return url, None, f"http_{error.code}"
    except URLError:
        return url, None, "url_error"
    except Exception:
        return url, None, "fetch_error"


def download_image(url: str, slug: str) -> Optional[str]:
    if not url:
        return None
    try:
        parsed = urlsplit(url)
        ext = Path(parsed.path).suffix.lower()
        if ext not in {".jpg", ".jpeg", ".png", ".webp"}:
            ext = ".jpg"
        target = IMG_PRODUCTS_DIR / f"{slug}{ext}"
        if target.exists():
            return f"img/products/{target.name}"
        req = Request(url, headers={"User-Agent": USER_AGENT})
        with urlopen(req, timeout=20) as resp:
            content = resp.read()
        if not content:
            return None
        IMG_PRODUCTS_DIR.mkdir(parents=True, exist_ok=True)
        target.write_bytes(content)
        return f"img/products/{target.name}"
    except Exception:
        return None


def build_existing_title_index(existing_products: Dict[str, dict]):
    index = {}
    entries = []
    for _, data in existing_products.items():
        title = str(data.get("title") or "").strip()
        if not title:
            continue
        key = normalize_for_match(title)
        if key and key not in index:
            index[key] = data
        entries.append((key, data))
    return index, entries


def build_file_index() -> Tuple[Dict[str, str], List[str]]:
    file_map = {}
    keys = []
    if not IMG_PRODUCTS_DIR.exists():
        return file_map, keys
    for file in IMG_PRODUCTS_DIR.iterdir():
        if not file.is_file():
            continue
        key = normalize_for_match(file.stem)
        file_map[key] = file.name
        keys.append(key)
    return file_map, keys


def find_existing_match(
    title: str,
    existing_index: Dict[str, dict],
    existing_entries: List[Tuple[str, dict]],
) -> Optional[dict]:
    candidates = [normalize_for_match(title), normalize_for_match(re.sub(r"\bx\s*1\s*und\b", "", title, flags=re.I))]
    for key in candidates:
        if key and key in existing_index:
            return existing_index[key]

    for key in candidates:
        if not key:
            continue
        close = difflib.get_close_matches(key, list(existing_index.keys()), n=1, cutoff=0.92)
        if close:
            return existing_index[close[0]]

    best = None
    best_ratio = 0.0
    for key in candidates:
        if not key:
            continue
        for ek, data in existing_entries:
            if not ek:
                continue
            ratio = difflib.SequenceMatcher(None, key, ek).ratio()
            if ratio > best_ratio:
                best_ratio = ratio
                best = data
    if best_ratio >= 0.94:
        return best
    return None


def find_image_for_product(
    title: str,
    slug: str,
    existing_match: Optional[dict],
    file_map: Dict[str, str],
    file_keys: List[str],
    fetched_data: Optional[dict],
) -> str:
    # 1) Match by local filenames (exact or close)
    candidate_keys = [
        normalize_for_match(title),
        normalize_for_match(re.sub(r"\bx\s*1\s*und\b", "", title, flags=re.I)),
        normalize_for_match(slug.replace("-", " ")),
    ]
    for key in candidate_keys:
        if key in file_map:
            return f"img/products/{file_map[key]}"

    for key in candidate_keys:
        close = difflib.get_close_matches(key, file_keys, n=1, cutoff=0.84)
        if close:
            return f"img/products/{file_map[close[0]]}"

    # 2) Download from source if available
    image_url = (fetched_data or {}).get("image_url") if fetched_data else ""
    downloaded = download_image(image_url, slug)
    if downloaded:
        return downloaded

    # 3) Reuse image path from previous catalog as fallback
    if existing_match:
        current = str(existing_match.get("image") or "").strip()
        current = current.replace("%20", " ")
        if current.startswith("assets/img/"):
            current = current.replace("assets/img/", "img/")
        if current and (ROOT / current).exists():
            return current

    return "img/icon.svg"


def has_title_mismatch(title: str, fetched_name: str) -> bool:
    base_title = normalize_for_match(title)
    base_fetched = normalize_for_match(fetched_name)
    if not base_fetched:
        return False

    ratio = difflib.SequenceMatcher(None, base_title, base_fetched).ratio()
    if ratio < 0.82:
        return True

    title_colors = set(c for c in extract_colors(title) if c != "Clasico")
    fetched_colors = set(c for c in extract_colors(fetched_name) if c != "Clasico")
    if title_colors and fetched_colors and not (title_colors & fetched_colors):
        return True
    return False


def build_product_record(
    row: dict,
    existing_match: Optional[dict],
    fetched_data: Optional[dict],
    file_map: Dict[str, str],
    file_keys: List[str],
) -> dict:
    title = row["title"].strip()
    slug = slugify(title)

    if fetched_data and has_title_mismatch(title, fetched_data.get("name", "")):
        fetched_data = None

    category = categorize_title(title, (fetched_data or {}).get("categories"))

    source_short = clean_html_to_text((fetched_data or {}).get("short_description_html", ""))
    source_desc = clean_html_to_text((fetched_data or {}).get("description_html", ""))
    source_meta = clean_html_to_text((fetched_data or {}).get("meta_description", ""))

    if not source_short and existing_match:
        source_short = str(existing_match.get("description") or "")
    if not source_desc and existing_match:
        source_desc = str(existing_match.get("longDescription") or existing_match.get("description") or "")
    if not source_meta:
        source_meta = source_short

    short_description = first_sentences(source_short or source_meta, count=1)
    if len(short_description) < 20:
        short_description = f"Producto de {category} para uso diario en estudio, oficina y hogar."

    description = first_sentences(source_meta or source_desc or short_description, count=1)
    if len(description) < 30:
        description = f"{title} disponible en Libreria Belen con calidad confiable para uso frecuente."

    detail_text = first_sentences(source_desc or source_meta or description, count=3)
    brand = (fetched_data or {}).get("brand") or extract_brand_from_title(title) or "Generica"
    specs = extract_specs(title)
    if specs:
        detail_text = f"{detail_text} {' '.join(specs)}"
    if len(detail_text) < 80:
        detail_text = f"{detail_text} Ideal para actividades escolares, oficina y proyectos personales."

    usage = build_usage(title, category)
    if existing_match and not usage:
        existing_usage = existing_match.get("usage")
        if isinstance(existing_usage, list) and existing_usage:
            usage = existing_usage[:3]

    details = [f"Marca: {brand}.", f"Categoria: {category}."]
    for spec in specs:
        if len(details) >= 3:
            break
        details.append(spec)
    details.append("Producto con precio actualizado para venta inmediata.")

    unit_price = row.get("price") or safe_float((fetched_data or {}).get("price")) or 0.0
    if unit_price <= 0 and existing_match:
        unit_price = safe_float(existing_match.get("price")) or 0.0
    unit_price = round(float(unit_price), 2)

    image = find_image_for_product(title, slug, existing_match, file_map, file_keys, fetched_data)

    tag = "oferta" if unit_price <= 3.5 else "nuevo"
    link = row.get("link", "")
    rating = safe_float((existing_match or {}).get("rating")) or 4.5

    return {
        "slug": slug,
        "title": title,
        "short": short_description,
        "description": description,
        "longDescription": detail_text,
        "price": unit_price,
        "rating": round(float(rating), 1),
        "stock": 20,
        "tag": tag,
        "category": category,
        "brand": brand,
        "link": link,
        "colors": extract_colors(title),
        "sizes": extract_sizes(title),
        "details": details,
        "usage": usage,
        "reviews": [
            {"name": "Cliente", "text": "Buen producto y entrega rapida."},
            {"name": "Libreria Belen", "text": "Precio actualizado y stock disponible."},
        ],
        "image": image,
        "gallery": [],
    }


def write_products_js(products: Dict[str, dict], offers: List[dict]):
    content = []
    content.append("(() => {")
    content.append("  const PRODUCTS = " + json.dumps(products, ensure_ascii=True, indent=2) + ";")
    content.append("  const OFFERS = " + json.dumps(offers, ensure_ascii=True, indent=2) + ";")
    content.append("  window.PRODUCTS = PRODUCTS;")
    content.append("  window.OFFERS = OFFERS;")
    content.append("})();")
    PRODUCTS_JS.write_text("\n".join(content) + "\n", encoding="utf-8")


def clean_unused_images(products: Dict[str, dict]) -> Tuple[int, int]:
    referenced = set()
    for item in products.values():
        image = str(item.get("image") or "")
        if image.startswith("img/products/"):
            referenced.add(image.split("/", 2)[-1])

    total = 0
    removed = 0
    if IMG_PRODUCTS_DIR.exists():
        for file in IMG_PRODUCTS_DIR.iterdir():
            if not file.is_file():
                continue
            total += 1
            if file.name not in referenced:
                file.unlink(missing_ok=True)
                removed += 1
    return total, removed


def main():
    rows = parse_sheet_rows()
    existing_products = load_existing_products(PRODUCTS_JS)
    existing_index, existing_entries = build_existing_title_index(existing_products)
    file_map, file_keys = build_file_index()

    unique_links = sorted({row["link"] for row in rows if row.get("link")})
    fetched_by_url: Dict[str, Optional[dict]] = {}
    fetch_errors: Dict[str, int] = {}

    with concurrent.futures.ThreadPoolExecutor(max_workers=10) as pool:
        futures = [pool.submit(fetch_product_info, url) for url in unique_links]
        for future in concurrent.futures.as_completed(futures):
            url, data, error = future.result()
            fetched_by_url[url] = data
            if error:
                fetch_errors[error] = fetch_errors.get(error, 0) + 1

    products_obj: Dict[str, dict] = {}
    slug_count = {}

    for row in rows:
        existing_match = find_existing_match(row["title"], existing_index, existing_entries)
        fetched_data = fetched_by_url.get(row.get("link", ""))
        record = build_product_record(row, existing_match, fetched_data, file_map, file_keys)

        base_slug = record.pop("slug")
        slug_count[base_slug] = slug_count.get(base_slug, 0) + 1
        slug = base_slug if slug_count[base_slug] == 1 else f"{base_slug}-{slug_count[base_slug]}"
        products_obj[slug] = record

    sorted_items = sorted(products_obj.items(), key=lambda x: x[1]["price"])
    offers = [
        {
            "title": item["title"],
            "description": item["description"],
            "price": f"S/ {item['price']:.2f}",
        }
        for _, item in sorted_items[:3]
    ]

    write_products_js(products_obj, offers)
    total_before, removed = clean_unused_images(products_obj)

    print(f"Rows processed: {len(rows)}")
    print(f"Products generated: {len(products_obj)}")
    print(f"Links fetched: {len(unique_links)}")
    print(f"Fetch error summary: {fetch_errors or 'none'}")
    print(f"Images before cleanup: {total_before}")
    print(f"Images removed as unused: {removed}")


if __name__ == "__main__":
    main()
